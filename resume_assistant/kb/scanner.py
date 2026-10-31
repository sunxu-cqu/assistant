"""知识库扫描：把「一个项目一个文件夹」的资料读成纯文本。

真实资料比测试资料脏得多，所以这里的重点不是「支持多少格式」，而是
**读不到的时候要说清楚为什么**：

- 中文 txt 大量是 GBK 编码，直接按 utf-8 读会静默变成乱码或空字符串
- 老版 .doc 不是 zip 包，python-docx 打不开，必须提示另存为 docx
- 扫描件 PDF 没有文字层，抽出来是空的，看起来像「资料不存在」
- 截图（png/jpg）同理

`scan_folder()` 把读不到的每个文件连同原因记进 `skipped`，由 `kb scan`
汇总给用户看——比让人对着「0 条证据」猜要强。
"""

from __future__ import annotations

import codecs
import hashlib
from dataclasses import dataclass, field
from pathlib import Path

from ..config import PROJECTS_DIR, settings

TEXT_SUFFIXES = {".md", ".markdown", ".txt", ".rst", ".csv", ".json", ".yaml", ".yml", ".log"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".webp", ".tif", ".tiff", ".heic"}
ARCHIVE_SUFFIXES = {".zip", ".rar", ".7z", ".tar", ".gz"}
SKIP_DIRS = {".git", "__pycache__", ".venv", "node_modules", ".idea", ".vscode"}
MAX_PDF_PAGES = 40
MAX_SHEET_ROWS = 300
MAX_FILES_PER_FOLDER = 80
# 单个文件最多读这么多字。项目文件夹里混进训练日志、导出的 csv 很常见，
# 一个 6MB 的 log 就能把整个项目的字符预算吃光，真正有内容的笔记反而被截断掉
MAX_CHARS_PER_FILE = 8000
# 单个文件超过这个大小直接跳过：实验数据集、录屏、模型权重都会塞进项目文件夹，
# 读进来既慢又是噪声（模型看到 5 万行 CSV 只会晕）
MAX_FILE_MB = 20


@dataclass
class ProjectFolder:
    name: str
    path: Path
    text: str = ""
    files: list[str] = field(default_factory=list)
    # [{"file": "报告.doc", "reason": "老版 .doc 打不开，请另存为 .docx"}]
    skipped: list[dict[str, str]] = field(default_factory=list)
    truncated: bool = False
    encodings: dict[str, str] = field(default_factory=dict)

    @property
    def fingerprint(self) -> str:
        h = hashlib.sha1()
        for p in sorted(self.path.rglob("*")):
            if p.is_file() and not any(part in SKIP_DIRS for part in p.parts):
                h.update(p.name.encode("utf-8", "ignore"))
                h.update(str(int(p.stat().st_mtime)).encode())
                h.update(str(p.stat().st_size).encode())
        return h.hexdigest()[:16]

    @property
    def chars(self) -> int:
        return len(self.text)


def decode_bytes(raw: bytes) -> tuple[str, str]:
    """按 BOM → utf-8 → gb18030 → big5 的顺序试，返回 (文本, 实际编码)。

    gb18030 是 GBK/GB2312 的超集，中文 Windows 上导出的 txt、csv 基本都是它。
    直接 utf-8 + errors="ignore" 会把整篇中文吃掉，而且不报错——这是最坑的一种。
    """
    if raw.startswith(codecs.BOM_UTF8):
        return raw.decode("utf-8-sig", "replace"), "utf-8-sig"
    if raw.startswith(codecs.BOM_UTF16_LE) or raw.startswith(codecs.BOM_UTF16_BE):
        return raw.decode("utf-16", "replace"), "utf-16"
    for enc in ("utf-8", "gb18030", "big5"):
        try:
            return raw.decode(enc), enc
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "replace"), "utf-8(有损)"


def _read_text_file(path: Path) -> tuple[str, str, str]:
    raw = path.read_bytes()
    if not raw.strip():
        return "", "", "文件是空的"
    text, enc = decode_bytes(raw)
    # 控制字符太多说明猜错编码了（正常中文文本不会有大量 \x00-\x08）
    if text and sum(1 for ch in text[:2000] if ord(ch) < 9) > len(text[:2000]) * 0.05:
        return "", "", "编码无法识别，读出来是乱码"
    return text, enc, ""


def _read_pdf(path: Path) -> tuple[str, str]:
    try:
        import pdfplumber
    except ImportError:
        return "", "缺少 pdfplumber，无法读 PDF"
    out: list[str] = []
    try:
        with pdfplumber.open(str(path)) as pdf:
            for page in pdf.pages[:MAX_PDF_PAGES]:
                out.append(page.extract_text() or "")
    except Exception as exc:  # noqa: BLE001
        return "", f"PDF 打开失败：{type(exc).__name__}"
    text = "\n".join(out)
    if not text.strip():
        return "", "PDF 没有文字层（扫描件/图片版），P0 不支持 OCR"
    return text, ""


def _read_docx(path: Path) -> tuple[str, str]:
    try:
        import docx
    except ImportError:
        return "", "缺少 python-docx，无法读 Word"
    try:
        doc = docx.Document(str(path))
    except Exception:  # noqa: BLE001
        if path.suffix.lower() == ".doc":
            return "", "老版 .doc 打不开，请用 Word 另存为 .docx 后重试"
        return "", "Word 文件损坏或格式异常"
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts), ""


def _read_pptx(path: Path) -> tuple[str, str]:
    try:
        from pptx import Presentation
    except ImportError:
        return "", "缺少 python-pptx，无法读 PPT"
    try:
        prs = Presentation(str(path))
    except Exception:  # noqa: BLE001
        return "", "PPT 文件损坏或格式异常"
    out: list[str] = []
    for i, slide in enumerate(prs.slides, 1):
        texts: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                t = (shape.text_frame.text or "").strip()
                if t:
                    texts.append(t)
            if getattr(shape, "has_table", False):
                try:
                    for row in shape.table.rows:
                        cells = [c.text.strip() for c in row.cells if c.text.strip()]
                        if cells:
                            texts.append(" | ".join(cells))
                except Exception:  # noqa: BLE001
                    pass
        if texts:
            out.append(f"[第 {i} 页]\n" + "\n".join(texts))
    if not out:
        return "", "PPT 里只有图片，没有可读文字"
    return "\n".join(out), ""


def _read_xlsx(path: Path) -> tuple[str, str]:
    try:
        import openpyxl
    except ImportError:
        return "", "缺少 openpyxl，无法读 Excel"
    try:
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001
        return "", "Excel 文件损坏或格式异常"
    out: list[str] = []
    try:
        for ws in wb.worksheets:
            rows: list[str] = []
            for n, row in enumerate(ws.iter_rows(values_only=True)):
                if n >= MAX_SHEET_ROWS:
                    rows.append(f"[超过 {MAX_SHEET_ROWS} 行，已截断]")
                    break
                cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if cells:
                    rows.append(" | ".join(cells))
            if rows:
                out.append(f"[工作表 {ws.title}]\n" + "\n".join(rows))
    finally:
        wb.close()
    if not out:
        return "", "Excel 是空的"
    return "\n".join(out), ""


def read_file(path: Path) -> tuple[str, str, str]:
    """返回 (文本, 编码, 跳过原因)。跳过原因非空表示这个文件没读到内容。"""
    suffix = path.suffix.lower()
    try:
        size_mb = path.stat().st_size / (1024 * 1024)
        if size_mb > MAX_FILE_MB:
            return "", "", (
                f"文件 {size_mb:.0f}MB，超过 {MAX_FILE_MB}MB 上限，"
                "多半是数据集/录屏，读进来也是噪声；有结论请写进 md"
            )
    except OSError as exc:
        return "", "", f"无法读取文件信息：{type(exc).__name__}"
    try:
        if suffix in TEXT_SUFFIXES:
            return _read_text_file(path)
        if suffix == ".pdf":
            text, why = _read_pdf(path)
            return text, "", why
        if suffix in {".docx", ".doc"}:
            text, why = _read_docx(path)
            return text, "", why
        if suffix == ".pptx":
            text, why = _read_pptx(path)
            return text, "", why
        if suffix in {".xlsx", ".xlsm"}:
            text, why = _read_xlsx(path)
            return text, "", why
        if suffix == ".ppt":
            return "", "", "老版 .ppt 打不开，请另存为 .pptx 后重试"
        if suffix in {".xls", ".et"}:
            return "", "", "老版 .xls 打不开，请另存为 .xlsx 后重试"
        if suffix in IMAGE_SUFFIXES:
            return "", "", "图片没有文字层，P0 不支持 OCR（重要内容请手动写进 md）"
        if suffix in ARCHIVE_SUFFIXES:
            return "", "", "压缩包不解压，请解压后放入文件夹"
    except Exception as exc:  # noqa: BLE001
        return "", "", f"读取失败：{type(exc).__name__}"
    return "", "", f"不支持的格式 {suffix or '(无扩展名)'}"


def _rank(rel: str) -> int:
    """决定先读哪个文件。

    一个项目文件夹里常有几十个代码文件 + 一份 README。截断是按字符数从后往前砍的，
    如果按文件名排序，README 很可能排在中间被砍掉，留下的全是代码。
    所以先读「叙述性」文件：README / 说明 / 报告 / 总结 / md。
    """
    low = rel.lower()
    name = low.rsplit("/", 1)[-1]
    if any(k in name for k in ("readme", "说明", "报告", "总结", "简介", "介绍", "方案", "记录")):
        return 0
    if low.endswith(".md"):
        return 0
    if low.endswith((".docx", ".doc", ".pdf", ".pptx", ".txt")):
        return 1
    return 2


def scan_folder(folder: Path) -> ProjectFolder:
    pf = ProjectFolder(name=folder.name, path=folder)
    candidates: list[tuple[int, str, Path]] = []
    for p in sorted(folder.rglob("*")):
        if not p.is_file():
            continue
        if any(part in SKIP_DIRS for part in p.parts):
            continue
        if p.name.startswith("~$"):  # Office 打开文件时留下的临时锁文件
            continue
        rel = str(p.relative_to(folder))
        candidates.append((_rank(rel), rel, p))

    candidates.sort(key=lambda x: (x[0], x[1]))
    if len(candidates) > MAX_FILES_PER_FOLDER:
        for _, rel, _p in candidates[MAX_FILES_PER_FOLDER:]:
            pf.skipped.append(
                {"file": rel, "reason": f"单个文件夹文件过多，只读前 {MAX_FILES_PER_FOLDER} 个"}
            )
        candidates = candidates[:MAX_FILES_PER_FOLDER]

    chunks: list[str] = []
    for _, rel, p in candidates:
        text, enc, why = read_file(p)
        if text.strip():
            pf.files.append(rel)
            if enc:
                pf.encodings[rel] = enc
            body = text.strip()
            if len(body) > MAX_CHARS_PER_FILE:
                body = body[:MAX_CHARS_PER_FILE] + f"\n[本文件超过 {MAX_CHARS_PER_FILE} 字，已截断]"
            chunks.append(f"\n\n===== 文件：{rel} =====\n{body}")
        else:
            pf.skipped.append({"file": rel, "reason": why or "没有读到文字"})

    full = "".join(chunks).strip()
    if len(full) > settings.max_project_chars:
        full = full[: settings.max_project_chars] + "\n\n[内容已截断]"
        pf.truncated = True
    pf.text = full
    return pf


def scan_all(base: Path | None = None) -> list[ProjectFolder]:
    base = base or PROJECTS_DIR
    if not base.exists():
        return []
    out: list[ProjectFolder] = []
    for child in sorted(base.iterdir()):
        if not child.is_dir() or child.name in SKIP_DIRS or child.name.startswith("."):
            continue
        out.append(scan_folder(child))
    return out
